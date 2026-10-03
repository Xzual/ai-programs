from pathlib import Path
from collections import Counter
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = Path(r"C:\Users\arday\Desktop\ai programs\EDITH_154_Ozellik_Durum_Raporu.docx")


def rows(category, entries):
    return [(category, *entry) for entry in entries]


DATA = []

DATA += rows("Ürün vizyonu ve masaüstü operatörü", [
    (1, "Ses -> niyet -> skill -> plan -> observe -> act -> verify -> sonuç ürün akışı", "PARTIAL", "Chat 4, 5, 6", "Skill/task/verification sözleşmeleri var; gerçek genel operatör zinciri bağlı değil."),
    (2, "Desktop ana yüzeyleri Voice Room, Computer Use ve Crypto; teknik ekranlar Advanced altında", "DONE", "Chat 2", "Ana navigasyon ve odak yapısı uygulandı; responsive UI testleri geçti."),
    (3, "Mobil ana yüzeyleri Voice, Computers, Active Tasks, Files, Research, History, Quick Actions ve Settings", "PARTIAL", "Chat 2, 4", "Compose ekranları ve sözleşmeler var; tüm yüzeylerin fiziksel cihaz E2E kanıtı yok."),
    (4, "Android 13 ve üstü destek", "DONE", "Chat 2", "minSdk 33, target/compile 35; API34 üzerinde 4/4 instrumentation geçti."),
    (5, "Mobilin bağımsız AI değil, eşleştirilen Desktop EDITH'in güvenli istemcisi olması", "DONE", "Chat 4, 2", "Owner/device lineage, pairing ve remote-client mimarisi kod ve testlerle tanımlı."),
    (6, "İnsan gibi görünür mouse, click, type, scroll, highlight ve gerçek Computer Use", "BLOCKED", "Chat 6, 8", "Native capture/mouse/keyboard/UIA/OCR runtime bağlı değil; gerçek masaüstü kabul koşusu yok."),
    (7, "Operator durumları Observing, Planning, Locating, Moving, Clicking, Typing, Scrolling, Verifying, Success ve Error", "PARTIAL", "Chat 2, 6", "Durum modeli ve UI sunumu var; tamamı gerçek native operatörle doğrulanmadı."),
    (8, "Hedef bulma sırası UI Automation, accessibility, window-relative, OCR/vision ve coordinate fallback", "PARTIAL", "Chat 6, 8", "Fallback sözleşmeleri mevcut; gerçek UIA/OCR ve karma DPI koşusu yok."),
    (9, "Observe -> Locate -> Act -> Verify döngüsü", "PARTIAL", "Chat 4, 6", "Dispatch/verification ayrımı ve test harness'i var; görünür gerçek uygulama E2E eksik."),
    (10, "Kör tıklama engeli, sınırlı retry ve fail-closed", "PARTIAL", "Chat 6, 8", "Stale target ve bounded retry testleri geçti; gerçek hedef üzerinde kabul koşusu yok."),
    (11, "Universal Desktop Operator", "PARTIAL", "Chat 6", "Genel sözleşme/oturum güvenliği var; çalışan universal native operator yok."),
    (12, "Steam, Spotify, Discord, Explorer, Chrome/Edge, Notepad, Office, Settings ve medya App Profiles", "PARTIAL", "Chat 5, 6", "Profil/capability altyapısı var; app-specific runtime operatörleri tamamlanmadı."),
    (13, "Steam profili; arama, kütüphane, indirme, durum ve bandwidth", "BLOCKED", "Chat 6, 8", "Gerçek Steam operatörü ve güvenli kabul koşusu yok."),
    (14, "Spotify profili; arama, oynatma, pause, next/prev, volume ve doğrulama", "BLOCKED", "Chat 6, 8", "Spotify operator runtime kanıtı yok."),
    (15, "File Explorer profili; klasör, arama, taşıma/kopyalama ve undo", "NOT_STARTED", "Chat 6, 8", "File Explorer'a özgü çalışan operatör veya kabul testi bulunmadı."),
])

DATA += rows("Browser ve araştırma", [
    (16, "Visible Browser Operator", "PARTIAL", "Chat 5, 6, 8", "Search/read-only temeli var; görünür browser operator UNBOUND."),
    (17, "Playwright/DOM hedef bulup gerçek mouse'u elemana götürme", "PARTIAL", "Chat 5, 6", "DOM/accessibility hedefleme temeli var; native mouse ile birleşik E2E yok."),
    (18, "URL, search, click, type, scroll, tabs, extraction, form fill, download ve verify", "PARTIAL", "Chat 5, 8", "Search/extraction politikaları var; navigate/form/download runtime configuration_required."),
    (19, "Ekranı ele geçirmeyen Background Browser/Research Worker", "PARTIAL", "Chat 5, 4", "Research run/worker sözleşmeleri var; varsayılan live acquisition worker bağlı değil."),
    (20, "Fast Research", "PARTIAL", "Chat 5", "Bounded research akışı var; gerçek kaynaklarla ürün kabulü tamamlanmadı."),
    (21, "Deep Research; çok kaynak, çelişki ve güvenilirlik", "PARTIAL", "Chat 5", "Claim/citation/freshness yapısı var; canlı çok kaynaklı E2E eksik."),
    (22, "JS-heavy siteler için Browser Research", "PARTIAL", "Chat 5, 8", "Browser research altyapısı var; gerçek JS-heavy kaynak kabulü yok."),
    (23, "Research tool set; search, fetch, extract, compare, fact-check, summarize, citations ve save", "PARTIAL", "Chat 5", "Araç/sözleşme temeli mevcut; tüm zincir production worker ile doğrulanmadı."),
    (24, "Resmi, GitHub/teknik ve community/forum Parallel Research worker'ları", "PARTIAL", "Chat 5", "Paralel araştırma tasarımı/üreticileri var; canlı worker matrisi kanıtlanmadı."),
])

DATA += rows("Obsidian ve araştırma belleği", [
    (25, "Obsidian ilk açılış Vault klasörü seçimi", "DONE", "Chat 2, 5, 6, 8", "Paketli EXE first-run picker bağımsız kabul edildi."),
    (26, "Windows folder picker", "DONE", "Chat 6, 8", "Native picker gerçek paketli Windows akışında doğrulandı."),
    (27, "Vault path'in cihaz bazlı local config'e kaydı", "DONE", "Chat 5, 6", "Seçim restart sonrasında tekrar kullanıldı."),
    (28, "Vault'u tekrar sormama; yalnız Settings üzerinden değiştirme", "DONE", "Chat 2, 6, 8", "Restart reuse, change ve revoke akışları geçti."),
    (29, "Vault unavailable durumunda popup spam olmadan degraded/unavailable", "DONE", "Chat 2, 8", "DEGRADED ve cancel davranışları paketli akışta doğrulandı."),
    (30, "Hardcoded D:\\EDİTH\\EDİTH path yasağı", "DONE", "Chat 5", "Hardcoded kullanıcı yolu kaldırıldı; kullanıcı seçimi ve güvenli config kullanılıyor."),
    (31, "Unicode, Türkçe ve boşluklu path desteği", "DONE", "Chat 5, 8", "Unicode/containment/path testleri geçti."),
    (32, "Test Vault ile gerçek Vault izolasyonu", "DONE", "Chat 5, 6, 8", "Fixed vault ve .edith envanteri testlerden sonra değişmedi."),
    (33, "UserVaultProvider ve SandboxVaultProvider ayrımı", "DONE", "Chat 5", "Provider sınırı ve native selection credential izolasyonu test edildi."),
    (34, "EDITH_TEST_MODE güvenliği", "DONE", "Chat 5, 6", "Test-mode/sandbox sınırları ve reparse/symlink reddi geçti."),
    (35, "Research Journal günlük dosyası Research/Daily/YYYY-MM-DD.md", "PARTIAL", "Chat 5", "Journal provider/sandbox testi var; production journal provider bağlı değil."),
    (36, "Aynı gün tek journal ve timestamp bölümleriyle append", "PARTIAL", "Chat 5", "Yazma sözleşmesi mevcut; gerçek kullanıcı vault'unda production E2E yok."),
    (37, "Research Topics klasörü", "PARTIAL", "Chat 5", "Knowledge/research modeli var; gerçek vault içerik üretimi kabul edilmedi."),
    (38, "Research Index ve MOC", "PARTIAL", "Chat 5", "Index/graph temeli var; gerçek production vault zinciri eksik."),
    (39, "Daily, topic ve project/task arasında wikilink", "PARTIAL", "Chat 5", "İlişki modeli var; gerçek journal-to-project kabul kanıtı yok."),
    (40, "Research Memory ile yalnız yeni/değişen bilgiyi araştırma", "PARTIAL", "Chat 5", "Freshness/delta sözleşmeleri test edildi; canlı tekrar araştırma E2E yok."),
])

DATA += rows("Bağlam, otomasyon ve yerel yetenekler", [
    (41, "Aktif uygulama, pencere, dosya, tab, task ve recent context awareness", "PARTIAL", "Chat 6, 5", "Metadata-only context üreticileri var; tam OS/tab bağlamı runtime kabulü yok."),
    (42, "Doğal takip komutları; bunu kaydet, onu aç, bir öncekini gönder", "PARTIAL", "Chat 5, 4", "Referans/intent sözleşmeleri var; gerçek çok adımlı kullanıcı E2E eksik."),
    (43, "Kısa süreli Screen Memory", "PARTIAL", "Chat 6, 5", "Retention ve private-data dışlama testleri var; gerçek ekran target recovery eksik."),
    (44, "Task Continuation; dün bırakılan yerden devam", "PARTIAL", "Chat 5, 4", "Checkpoint/recovery temeli var; restart sonrası gerçek mission recovery yok."),
    (45, "Clipboard Intelligence ve PC-mobile clipboard", "PARTIAL", "Chat 4, 5", "TTL/no-persist/secret rejection sözleşmeleri geçti; gerçek cihaz teslimi yok."),
    (46, "Windows Notification Intelligence", "PARTIAL", "Chat 6, 5", "Metadata ve sınıflandırma temeli var; OS notification source bağlı değil."),
    (47, "Macro Recorder", "NOT_STARTED", "Chat 6, 8", "Recorder uygulaması ve testi bulunmadı."),
    (48, "Workflow Replay", "PARTIAL", "Chat 5, 6", "Playbook/replay sözleşme temeli var; kaydedilmiş gerçek işin modelsiz replay kabulü yok."),
    (49, "Self-healing workflow", "PARTIAL", "Chat 5, 6", "Stale/moved target ve bounded repair testleri var; gerçek değişmiş UI E2E yok."),
    (50, "Routine Learning ve otomasyon önerisi", "PARTIAL", "Chat 5", "Öneri/capability yaklaşımı var; tekrar algılama ürün kabulü yok."),
    (51, "Undo Everything ve undo manifest", "PARTIAL", "Chat 5, 6", "Reversible plan/cleanup sözleşmeleri var; kapsamlı gerçek undo geçmişi yok."),
    (52, "Semantic local File Intelligence", "PARTIAL", "Chat 5", "Scoped local search/provenance/redaction testleri geçti; gerçek kullanıcı dosyası E2E eksik."),
    (53, "Dosya, Obsidian, research, projects ve tasks için Unified Local Search", "PARTIAL", "Chat 5", "Birleşik search temeli var; tüm kaynaklarla production kabulü yok."),
    (54, "Smart Download Manager", "PARTIAL", "Chat 5, 6", "Byte progress/private-path redaction üreticisi var; browser download/watcher/mobile zinciri bağlı değil."),
    (55, "Meeting Mode", "NOT_STARTED", "Chat 2, 5", "Transcript-summary-action-items için tamamlanmış ürün/runtime kanıtı yok."),
    (56, "Focus Mode", "NOT_STARTED", "Chat 2, 5", "Focus Mode'a özgü çalışan OS bildirim politikası kanıtı yok."),
    (57, "Spotify, YouTube, sistem medyası ve volume Media Control", "PARTIAL", "Chat 6", "Media/device sözleşme temeli var; gerçek Spotify/YouTube operatörü yok."),
    (58, "Ses çıkışı, mic, Bluetooth ve brightness Device Control", "PARTIAL", "Chat 6, 4", "Güvenli metadata/capability temeli var; gerçek OS cihaz kontrol E2E yok."),
    (59, "UI'yi dondurmayan Background Jobs", "PARTIAL", "Chat 4, 5, 6", "Sidecar ve bounded job sözleşmeleri var; tüm research/index/transfer runtime'ı bağlı değil."),
    (60, "Multi-monitor koordinat ve aktif monitor desteği", "BLOCKED", "Chat 6, 8", "Native capability ve fiziksel multi-monitor karma DPI testi yok."),
    (61, "Quick Voice Commands; dur, geri al, sus, iptal et, müziği azalt", "PARTIAL", "Chat 2, 4", "STOP/cancel sözleşmeleri var; gerçek mikrofon ve tüm komutların E2E kanıtı yok."),
    (62, "Model/provider/skill'den bağımsız Personality Layer", "PARTIAL", "Chat 2, 5", "Provider/model/persona ayrımı yapıldı; tam JARVIS konuşma katmanı kabulü yok."),
    (63, "Safe Offline Mode", "PARTIAL", "Chat 5, 6", "Local/fail-closed parçalar var; birleşik offline model+skill+vault+operator E2E yok."),
    (64, "Proactive Suggestions, kendi başına başlatmama", "PARTIAL", "Chat 5, 4", "Safety sınırı mevcut; rutin öneri davranışı ürün kabulü yok."),
])

DATA += rows("Mobil ve çapraz cihaz", [
    (65, "Android 13+ ayrı EDITH Mobile uygulaması", "DONE", "Chat 2, 4", "Kotlin/Compose uygulaması build oluyor; API34 instrumentation geçti."),
    (66, "QR veya tek kullanımlık pairing code", "PARTIAL", "Chat 4, 2", "Kriptografik pairing sözleşmeleri/testleri var; fiziksel cihaz pairing yapılmadı."),
    (67, "Trusted device kimliği, revoke, reconnect ve session expiry", "PARTIAL", "Chat 4", "Authority/revoke/replay testleri geçti; fiziksel interoperability yok."),
    (68, "Birden fazla PC'ye hazırlık ve task'ın tek PC'ye bağlı olması", "PARTIAL", "Chat 4, 2", "Target computer/lineage sözleşmesi var; gerçek multi-PC testi yok."),
    (69, "Mobil premium dark blue/cyan minimal tasarım", "DONE", "Chat 2", "Compose tasarım ve responsive viewport kontrolleri mevcut."),
    (70, "Mobil ana menüler Voice, Computers, Tasks, Files ve Research", "DONE", "Chat 2", "Ana mobil navigasyon yüzeyleri uygulandı."),
    (71, "Voice ekranı; core, bağlantı, seçili PC, task ve transcript", "PARTIAL", "Chat 2", "UI var; gerçek mikrofon/Gemini Live ve desktop yanıt zinciri yok."),
    (72, "Mobil remote Steam, Spotify, research, file search ve app open komutları", "PARTIAL", "Chat 2, 4, 6", "Remote command sözleşmeleri var; Steam/Spotify/native operator bağlı değil."),
    (73, "Gelenler klasörü ve PDF/Images/Documents/Other ayrımı", "PARTIAL", "Chat 2, 4", "UI/metadata reducer temeli var; fiziksel dosya teslimi kabul edilmedi."),
    (74, "PC -> Mobile encrypted file transfer ve integrity", "PARTIAL", "Chat 4, 2", "Chunk/resume/SHA-256 sözleşmeleri geçti; fiziksel byte transferi yok."),
    (75, "Mobile -> PC file transfer", "PARTIAL", "Chat 4, 2", "Yön/handle sözleşmeleri var; gerçek cihazdan PC'ye transfer yok."),
    (76, "Research result actions; ses, telefonda aç, TXT, desktop ve Obsidian", "PARTIAL", "Chat 2, 4, 5", "Result-card aksiyon modeli var; open/send/read-aloud gerçek E2E değil."),
    (77, "Canonical task runtime", "DONE", "Chat 4, 5", "TaskId/origin/target/skill/state/progress contract testleri geçti."),
    (78, "Task states queued, running, waiting, verifying, completed, failed ve stopped", "DONE", "Chat 4, 2", "Canonical state modeli desktop/mobile sözleşmelerinde mevcut."),
    (79, "Gerçek progress ve fake yüzde yasağı", "DONE", "Chat 4, 5", "Progress provenance ve no-fake kuralları test edildi."),
    (80, "Progress modes percent, steps, bytes ve indeterminate", "DONE", "Chat 4, 2", "Canonical progress modları contract/UI'da mevcut."),
    (81, "Steam gerçek download yüzdesi, hız ve kalan veri", "BLOCKED", "Chat 6", "Steam runtime operatörü bağlı değil."),
    (82, "Research source progress", "PARTIAL", "Chat 5, 2", "Kaynak sayısı/progress contract'ı var; canlı worker progress kabulü yok."),
    (83, "File transfer gerçek byte progress", "PARTIAL", "Chat 4, 2", "Byte progress harness'i geçti; fiziksel transfer E2E yok."),
    (84, "Computer Use live status", "PARTIAL", "Chat 6, 2", "Durum eventleri/UI var; gerçek native operatör akışı yok."),
    (85, "Realtime WebSocket veya WebRTC data channel", "PARTIAL", "Chat 4", "Canonical realtime/reconnect temeli var; gerçek iki cihaz delivery yok."),
    (86, "Event ordering ve stale event koruması", "DONE", "Chat 4, 2, 8", "Sequence/replay/stale reducer testleri ve legacy quarantine geçti."),
    (87, "Push notifications", "BLOCKED", "Chat 4, 8", "Provider/credential lifecycle ve gerçek delivery yok."),
    (88, "Active Tasks ekranı", "DONE", "Chat 2", "Task ekranı ve canonical status sunumu var."),
    (89, "Task History ekranı", "DONE", "Chat 2", "History/timeline yüzeyi mevcut."),
    (90, "Mobile STOP ve Emergency Stop", "PARTIAL", "Chat 2, 4", "UI, authority ve exclusion testleri var; fiziksel cihaz E2E yok."),
    (91, "Live Remote View", "BLOCKED", "Chat 4, 6, 8", "Pull-only/privacy sözleşmesi var; gerçek frame capture ve delivery yok."),
    (92, "Ghost Cursor ve Presence Mode", "PARTIAL", "Chat 2, 6", "HUD/highlight/cursor sözleşme ve UI temeli var; native overlay eksik."),
    (93, "Mission Mode", "PARTIAL", "Chat 4, 5, 2", "Mission/task contract ve UI var; birleşik gerçek execution zinciri yok."),
    (94, "Task Chaining", "PARTIAL", "Chat 4, 5", "ResultRef/task contract temeli var; gerçek zincir kabulü yok."),
    (95, "Cross-device clipboard", "PARTIAL", "Chat 4, 2", "Güvenli clipboard sözleşmeleri var; gerçek cihaz delivery yok."),
    (96, "Smart PC Status", "PARTIAL", "Chat 4, 2", "PC status modeli var; bazı metrikler/native freshness gerçek runtime'da eksik."),
    (97, "Wake-on-LAN ve Ready bildirimi", "PARTIAL", "Chat 4, 6", "WOL/capability temeli var; fiziksel WOL koşusu ve push bildirimi yok."),
    (98, "Remote File Fetch", "PARTIAL", "Chat 4, 5", "Search/handle/transfer contract'ı var; gerçek PDF fetch E2E yok."),
    (99, "Remote File Drop", "PARTIAL", "Chat 4, 2", "Transfer temeli var; telefondan PC'ye gerçek dosya yok."),
    (100, "Event-Based Missions", "PARTIAL", "Chat 4, 5", "Watcher/mission contract'ı var; OS event source bağlı değil."),
    (101, "Time-Based Missions", "PARTIAL", "Chat 4, 5", "Typed mission/schedule temeli var; native background scheduler yok."),
    (102, "Offline Mobile Queue", "PARTIAL", "Chat 4, 2", "Encrypted queue/idempotency/reconnect testleri var; gerçek offline cihaz E2E yok."),
    (103, "Activity Timeline", "PARTIAL", "Chat 2, 4", "Canonical event timeline UI var; gerçek tüm adımların cihazlar arası kabulü yok."),
    (104, "Daily Timeline", "PARTIAL", "Chat 2, 5", "Activity/history temeli var; günlük birleşik task/research/file kabulü yok."),
    (105, "Voice Notes to Obsidian", "PARTIAL", "Chat 2, 5", "Audio/note ve vault sözleşmeleri var; gerçek capture-to-vault E2E yok."),
    (106, "Ask My Computer", "PARTIAL", "Chat 5, 2", "Local search var; gerçek dosyada open/send/summarize kabulü yok."),
    (107, "Smart Handoff", "PARTIAL", "Chat 4, 2", "Cross-device result/task sözleşmeleri var; gerçek cihaz handoff yok."),
    (108, "Camera Assist geleceğe hazırlık", "NOT_STARTED", "Chat 2, 4", "Kamera tabanlı yardım için uygulanmış runtime kanıtı yok."),
    (109, "Mobile Quick Actions ve Widget", "PARTIAL", "Chat 2", "Quick Action UI var; Android widget ve tüm aksiyonların runtime kabulü yok."),
    (110, "Multi-PC pairing", "PARTIAL", "Chat 4, 2", "Sözleşme hazırlığı var; gerçek birden fazla PC pairing testi yok."),
])

DATA += rows("Güvenlik, servis ve platform", [
    (111, "Hassas uygulamalarda screenshot, stream ve log sınırlayan Privacy Mode", "DONE", "Chat 4, 6, 8", "Denylist, redaction ve metadata-only politikaları test edildi."),
    (112, "Normal işlemlerde sürekli izin sormama", "PARTIAL", "Chat 4, 2", "Scoped owner/session izin modeli var; gerçek operatör akışında ergonomi E2E yok."),
    (113, "Owner mode ile düşük riskli explicit komutu doğrudan çalıştırma", "PARTIAL", "Chat 4", "Owner/auth/approval sınırı var; native low-risk action acceptance eksik."),
    (114, "Ödeme, satın alma, real trading, permanent delete, credential/security ve arbitrary shell sınırları", "DONE", "Chat 4, 7, 8", "Hard safety boundaries ve live-trading kilitleri test edildi."),
    (115, "Trusted Mobile Owner Session", "PARTIAL", "Chat 4", "Device/session/expiry modeli var; fiziksel pairing sonrası gerçek kullanım yok."),
    (116, "Remote auth, encryption, replay protection, device binding ve revoke", "DONE", "Chat 4, 8", "P-256/HKDF/AES-GCM, sequence/replay/tamper ve revocation testleri geçti."),
    (117, "Public unauthenticated control port yasağı", "DONE", "Chat 4", "Loopback default, owner auth, same-origin ve CSRF kontrolleri geçti."),
    (118, "Outbound secure relay veya tunnel yaklaşımı", "NOT_STARTED", "Chat 4", "Production relay/tunnel uygulaması ve acceptance kanıtı yok."),
    (119, "Mobil bağlantı için Background Desktop Service", "PARTIAL", "Chat 6, 4", "Managed sidecar/lifecycle var; sürekli remote service delivery tamamlanmadı."),
    (120, "Güvenli ve açık Auto-start opsiyonu", "PARTIAL", "Chat 6", "Lifecycle güvenliği var; kullanıcı kontrollü autostart acceptance kanıtı yok."),
    (121, "Android 13+ notification, background task, file picker ve battery-aware UX", "PARTIAL", "Chat 2, 4", "Android yapı/izin/UI temeli var; push/background/file-picker fiziksel E2E eksik."),
])

DATA += rows("Araştırma, kalite ve yönetişim", [
    (122, "UI-TARS, Browser Use, Playwright MCP, usecomputer, Windows CU MCP, Oya ve registry repo araştırması", "DONE", "Chat 5, 6", "Araştırma/architecture handoff'ları oluşturuldu ve tasarıma işlendi."),
    (123, "Clean-room adaptation ve lisans kontrolü", "DONE", "Chat 5, 6", "Doğrudan kör kopyalama yerine yerel sözleşme/adaptör yaklaşımı kullanıldı."),
    (124, "Skill Registry'nin Desktop Vision, UI Automation, Browser, Research, File, Mobile ve Task Runtime ile genişletilmesi", "PARTIAL", "Chat 5", "Canonical registry ve planning capability ayrımı var; bazı runtime adapter'ları blocked."),
    (125, "Canonical task data contract", "DONE", "Chat 4, 5, 8", "V2/V2.1 task/activity contract testleri geçti; drift bulunmadı."),
    (126, "Fake timer-progress yasağı ve progress derivation kuralları", "DONE", "Chat 4, 5", "Gerçek source/step/byte progress ve fail-closed parser testleri geçti."),
    (127, "Notepad, Calculator ve Explorer Desktop acceptance testleri", "BLOCKED", "Chat 6, 8", "Gerçek UIA/OCR native operator olmadığı için testler koşulmadı."),
    (128, "Search, click, result açma ve title verify Browser acceptance", "BLOCKED", "Chat 5, 8", "Görünür browser operator bağlı değil."),
    (129, "Steam safe testleri", "BLOCKED", "Chat 6, 8", "Steam operatörü yok."),
    (130, "Spotify/media safe testleri", "BLOCKED", "Chat 6, 8", "Spotify operatörü yok; yalnız genel safety sınırları var."),
    (131, "Background, progress, TXT save ve test vault Research testleri", "PARTIAL", "Chat 5, 8", "Contract/sandbox testleri var; canlı worker, TXT ve sesli özet zinciri eksik."),
    (132, "Pairing, task sync, STOP, transfer, push ve reconnect Mobile acceptance", "BLOCKED", "Chat 2, 4, 8", "API34 instrumentation geçti; fiziksel iki cihaz/push/transfer kabulü yok."),
    (133, "Obsidian picker, restart, Settings change ve test vault isolation acceptance", "DONE", "Chat 2, 5, 6, 8", "Paketli Phase 11 functional/security acceptance tamamlandı."),
    (134, "Android 13, 14 ve 15 test hedefi", "PARTIAL", "Chat 2, 8", "API34 geçti; API33 incomplete, API35 image var fakat runtime testi yok."),
    (135, "Voice, target, mouse, browser, realtime, notification ve file throughput performans ölçümleri", "PARTIAL", "Chat 6, 8", "Build/startup gözlemleri var; kapsamlı latency/CPU/memory/throughput bütçesi yok."),
    (136, "Replay'in ilk çalışmadan daha az model çağrısı kullanması", "NOT_STARTED", "Chat 5, 8", "Tam macro recorder/replay olmadığı için ölçüm kanıtı yok."),
    (137, "Local telemetry; task, skill, steps, retries, duration ve verification", "PARTIAL", "Chat 4, 5", "Event/diagnostics/history kayıtları var; tam yerel workflow history kabulü yok."),
    (138, "Secret logging yasağı", "DONE", "Chat 4, 7, 8", "Secret redaction, browser storage temizliği ve backend-only secret testleri geçti."),
    (139, "Chat owner dağılımı", "DONE", "Chat 0, 2, 4, 5, 6, 7, 8", "Sorumluluklar chat bazında ayrıldı ve handoff raporları üretildi."),
    (140, "Dependency waves", "DONE", "Chat 0", "Contract -> native/UI -> integration -> bağımsız QA sırası izlendi."),
    (141, "One-file-owner-at-a-time", "DONE", "Chat 0", "Ownership/dependency yaklaşımı raporlarda korundu; çapraz düzeltmeler yönlendirildi."),
    (142, "No-fake ve honest runtime", "DONE", "Tüm chatler, Chat 8", "Unavailable/configuration_required durumları dürüst gösteriliyor; fake READY/progress reddediliyor."),
    (143, "Mevcut EDITH güvenliklerini bozmama", "DONE", "Chat 4, 6, 7, 8", "Security/regression testleri ve bağımsız QA geçti."),
    (144, "Crypto'nun demo-only kalması ve real trading olmaması", "DONE", "Chat 7, 8", "10.000 CR demo ledger, HOLD/no-trade ve live/private order kilitleri doğrulandı."),
])

DATA += rows("Final senaryolar ve başarı tanımı", [
    (145, "Desktop Voice ile Steam final senaryosu", "BLOCKED", "Chat 2, 6, 8", "Wake word/gerçek voice ve Steam native operator yok."),
    (146, "Mobile remote Steam final senaryosu", "BLOCKED", "Chat 2, 4, 6, 8", "Fiziksel pairing ve Steam operator yok."),
    (147, "PDF bulup telefona gönderme final senaryosu", "BLOCKED", "Chat 4, 5, 8", "Search/transfer sözleşmesi var; fiziksel byte transferi yok."),
    (148, "Background research, TXT ve sesli özet final senaryosu", "BLOCKED", "Chat 5, 8", "Live acquisition worker ve gerçek audio/TXT zinciri bağlı değil."),
    (149, "Obsidian first-run final senaryosu", "DONE", "Chat 2, 5, 6, 8", "Paketli picker, restart reuse, change/revoke ve izolasyon kabul edildi."),
    (150, "PC'de başlayan task'ın mobilde realtime görünmesi", "BLOCKED", "Chat 2, 4, 8", "Canonical realtime var; gerçek iki cihazlı delivery yok."),
    (151, "Offline queue final senaryosu", "BLOCKED", "Chat 2, 4, 8", "Reducer/harness var; gerçek offline PC-mobile reconnect acceptance yok."),
    (152, "Ask My Computer final senaryosu", "PARTIAL", "Chat 5, 8", "Scoped search var; gerçek kullanıcı PDF open/send/summarize E2E yok."),
    (153, "Tüm modüller için DONE, PARTIAL, NOT_STARTED ve BLOCKED final raporu", "DONE", "Chat 0, 8", "Master raporlar ve bu 154 maddelik kanıt matrisi hazırlandı."),
    (154, "Konuş, PC işi yapsın, gerçek progress, transfer, Obsidian memory ve düşük izin spam'i başarı tanımı", "PARTIAL", "Tüm chatler", "Obsidian ve contract/security güçlü; gerçek operator, voice, transfer ve mobil E2E tamamlanmadı."),
])


assert len(DATA) == 154, len(DATA)
assert [r[1] for r in DATA] == list(range(1, 155))


STATUS_COLORS = {
    "DONE": "D9EAD3",
    "PARTIAL": "FFF2CC",
    "NOT_STARTED": "E7E6E6",
    "BLOCKED": "F4CCCC",
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_font(run, size=9, bold=False, color="000000"):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.62)
section.bottom_margin = Inches(0.62)
section.left_margin = Inches(0.62)
section.right_margin = Inches(0.62)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor(0, 0, 0)
for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12)):
    st = styles[style_name]
    st.font.name = "Aptos Display" if style_name != "Normal" else "Aptos"
    st._element.rPr.rFonts.set(qn("w:ascii"), st.font.name)
    st._element.rPr.rFonts.set(qn("w:hAnsi"), st.font.name)
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("EDITH 154 Özellik Durum Raporu")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("Chat 2  Chat 4  Chat 5  Chat 6  Chat 7  Chat 8 kanıt konsolidasyonu")
set_font(r, 11, False, "404040")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(8)
r = p.add_run(
    "Bu rapor, 154 maddelik EDITH ürün vizyonunu mevcut chat raporları, bağımsız QA sonuçları ve gerçek runtime kanıtlarıyla tek tek sınıflandırır. "
    "Bir sözleşmenin veya arayüzün bulunması, gerçek cihaz ya da native uçtan uca kabulü yoksa DONE sayılmamıştır."
)
set_font(r, 10.5)

counts = Counter(r[3] for r in DATA)
doc.add_heading("Yönetici Özeti", level=1)
summary_table = doc.add_table(rows=2, cols=5)
summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
summary_table.autofit = False
set_repeat_table_header(summary_table.rows[0])
headers = ["Toplam", "DONE", "PARTIAL", "NOT STARTED", "BLOCKED"]
values = [154, counts["DONE"], counts["PARTIAL"], counts["NOT_STARTED"], counts["BLOCKED"]]
for i, value in enumerate(headers):
    cell = summary_table.cell(0, i)
    set_cell_shading(cell, "1F4E78")
    cell.text = value
    for run in cell.paragraphs[0].runs:
        set_font(run, 9, True, "FFFFFF")
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
for i, value in enumerate(values):
    cell = summary_table.cell(1, i)
    cell.text = str(value)
    if i > 0:
        set_cell_shading(cell, STATUS_COLORS[headers[i].replace(" ", "_")])
    for run in cell.paragraphs[0].runs:
        set_font(run, 11, True)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
for row in summary_table.rows:
    for cell in row.cells:
        set_cell_margins(cell, 110, 90, 110, 90)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(9)
r = p.add_run("Nihai sonuç: ")
set_font(r, 10.5, True)
r = p.add_run(
    "Temel güvenlik, canonical contract, paketleme ve Obsidian first-run güçlü biçimde tamamlandı. "
    "Ancak gerçek Computer Use ve Browser operator, Steam/Spotify, fiziksel mobil, dosya transferi, push, live view ve ses zincirleri tamamlanmadığı için 154 maddelik vizyon bütünüyle bitmiş değildir."
)
set_font(r, 10.5)

doc.add_heading("Durum Tanımları", level=1)
for status, explanation in [
    ("DONE", "İstenen kapsam mevcut kanıtla tamamlanmış ve ilgili test veya runtime kabulü yapılmış."),
    ("PARTIAL", "Kod, arayüz, contract veya test temeli var; gerçek E2E ya da kapsamın bir bölümü eksik."),
    ("NOT STARTED", "Uygulanmış ve doğrulanmış kanıt bulunamadı."),
    ("BLOCKED", "Temel veya contract bulunabilir; gerekli native runtime, cihaz, provider ya da dış koşul olmadığı için kabul tamamlanamadı."),
]:
    p = doc.add_paragraph(style=None)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(status + "  ")
    set_font(r, 10, True)
    r = p.add_run(explanation)
    set_font(r, 10)

current_category = None
for category, num, requirement, status, chat, evidence in DATA:
    if category != current_category:
        current_category = category
        doc.add_heading(category, level=1)
        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = [Inches(0.38), Inches(2.57), Inches(0.95), Inches(3.18)]
        for cell, width in zip(table.rows[0].cells, widths):
            cell.width = width
        hdr = table.rows[0].cells
        for i, label in enumerate(["No", "Özellik", "Durum", "Chat ve kanıt"]):
            hdr[i].text = label
            set_cell_shading(hdr[i], "1F4E78")
            hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in hdr[i].paragraphs[0].runs:
                set_font(run, 9, True, "FFFFFF")
            set_cell_margins(hdr[i])
        set_repeat_table_header(table.rows[0])

    data_row = table.add_row()
    set_cant_split(data_row)
    cells = data_row.cells
    cells[0].text = str(num)
    cells[1].text = requirement
    cells[2].text = status.replace("_", " ")
    cells[3].text = ""
    p = cells[3].paragraphs[0]
    r = p.add_run(chat + " - ")
    set_font(r, 8.5, True)
    r = p.add_run(evidence)
    set_font(r, 8.5)
    set_cell_shading(cells[2], STATUS_COLORS[status])
    for idx, cell in enumerate(cells):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if idx != 3:
            for run in cell.paragraphs[0].runs:
                set_font(run, 8.5, idx == 2)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        cell.paragraphs[0].paragraph_format.line_spacing = 1.0
    cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Chat Bazlı Sonuç", level=1)
chat_results = [
    ("Chat 2", "UI, responsive yüzeyler, canonical status, mobile Compose ve API34 instrumentation alanlarında güçlü. API33/35 runtime ve fiziksel cihaz kabulü eksik."),
    ("Chat 4", "Owner/session, CSRF, server-side permission, pairing, realtime ve cross-device güvenlik sözleşmelerini tamamladı. Gerçek push/provider ve iki cihazlı teslim eksik."),
    ("Chat 5", "Registry, research/knowledge, local search ve Obsidian sözleşmelerini geliştirdi. Canlı browser/research worker ve bazı workflow özellikleri eksik."),
    ("Chat 6", "Tauri lifecycle, sidecar, single-instance, portable ve paketleme işlerini tamamladı. Gerçek UIA/OCR/visible operator ve app-specific operatörler eksik."),
    ("Chat 7", "Crypto demo-only güvenliğini, packaged runtime'ı ve legacy data temizliğini tamamladı."),
    ("Chat 8", "Bağımsız QA, Obsidian acceptance ve fresh release setini doğruladı; master ürünü native/device/E2E ve imza eksikleri nedeniyle NOT_READY tuttu."),
]
for name, result in chat_results:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(name + "  ")
    set_font(r, 10, True)
    r = p.add_run(result)
    set_font(r, 10)

doc.add_heading("Başlıca Kalan İşler", level=1)
remaining = [
    "Gerçek Windows UI Automation, accessibility, OCR ve görünür mouse/keyboard operator runtime'ı.",
    "Görünür Browser Operator ve Playwright/DOM hedeflerinin native mouse ile birleşik kabulü.",
    "Steam, Spotify ve File Explorer app-specific operatörleri.",
    "Fiziksel Android pairing, iki cihazlı realtime, STOP, transfer, push ve reconnect testleri.",
    "Gerçek mikrofon/Gemini Live, voice-to-skill ve sesli sonuç zinciri.",
    "Macro Recorder, tam Workflow Replay ve model-call azaltımı ölçümü.",
    "Multi-monitor, native overlay, live view, watcher ve OS notification kaynakları.",
    "Authenticode imzalı release seti ve temiz/ikinci Windows makinesi acceptance testleri.",
]
for item in remaining:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(item)
    set_font(r, 10)

doc.add_heading("Kullanılan Kanıtlar", level=1)
sources = [
    "docs/EDITH_MASTER_PRODUCT_EXPANSION_FINAL_REPORT_2026-09-28.md",
    "docs/EDITH_MASTER_RELEASE_BLOCKER_FINAL_2026-09-27.md",
    "docs/EDITH_CHAT8_COMPUTER_BROWSER_SKILL_UPGRADE_FINAL_AUDIT_2026-09-28.md",
    "docs/EDITH_CHAT8_PHASE9_FINAL_INDEPENDENT_QA_2026-09-28.md",
    "docs/EDITH_CHAT8_PHASE11H_FINAL_P1_REACCEPTANCE_2026-09-29.md",
    "docs/EDITH_CHAT8_PHASE12B_FRESH_RELEASE_SET_REACCEPTANCE_2026-09-29.md",
    "docs/EDITH_CHAT2_WAVE2_FRONTEND_REPORT_2026-09-27.md",
    "docs/EDITH_WAVE1_CHAT4_BACKEND_SECURITY_REPORT_2026-09-27.md",
    "docs/EDITH_CHAT5_WAVE2_REGISTRY_OBSIDIAN_REPORT_2026-09-27.md",
    "docs/EDITH_CHAT6_WAVE1_DESKTOP_RELEASE_REPORT_2026-09-27.md",
    "docs/EDITH_CHAT7_CRYPTO_PACKAGED_EVIDENCE_2026-09-27.md",
    "Chat 2, 4, 5, 6, 7 ve 8 son çalışma kayıtları",
]
for source in sources:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(source)
    set_font(r, 9)

for section in doc.sections:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("EDITH 154 Özellik Durum Raporu  |  29 Eylül 2026")
    set_font(r, 8, False, "666666")

doc.core_properties.title = "EDITH 154 Özellik Durum Raporu"
doc.core_properties.subject = "Chat bazlı özellik kabul matrisi"
doc.core_properties.author = "Codex"
doc.core_properties.keywords = "EDITH, özellik, durum, QA, Computer Use, Mobile, Obsidian"
doc.save(OUT)
print(OUT)
print(dict(counts))
